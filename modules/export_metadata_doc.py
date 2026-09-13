"""
Экспорт названия и описания каждого клипа: Excel (.xlsx) или Markdown (.md)
(удобно для Google Sheets / Excel без облачных ссылок).
"""

from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path


def _natural_clip_sort(paths: list[Path]) -> list[Path]:
    def key(p: Path) -> tuple[int, str]:
        m = re.search(r"clip_(\d+)", p.name, re.I)
        return (int(m.group(1)), p.name) if m else (0, p.name)

    return sorted(paths, key=key)


def _load_project_meta(project_dir: Path) -> dict:
    pj = project_dir / "project.json"
    if not pj.exists():
        return {}
    try:
        return json.loads(pj.read_text(encoding="utf-8")) or {}
    except Exception:
        return {}


def _resolve_meta_sources(
    project_dir: Path,
    clips_subdir: str,
) -> tuple[dict, list[Path]]:
    project_dir = Path(project_dir)
    final_dir = project_dir / clips_subdir
    if not final_dir.is_dir():
        raise FileNotFoundError(f"Нет папки: {final_dir}")

    meta_files = list(final_dir.glob("*_meta.json"))
    meta_files = _natural_clip_sort(meta_files)
    if not meta_files:
        raise FileNotFoundError(f"Нет *_meta.json в {final_dir}")

    return _load_project_meta(project_dir), meta_files


def _iter_clip_rows(
    project_dir: Path,
    clips_subdir: str,
    *,
    include_tags: bool,
) -> tuple[dict, list[tuple[str, str, str, str]]]:
    """
    Returns (project_json, rows) где каждая строка:
    (filename.mp4, title, description, tags_csv или "")
    """
    proj, meta_files = _resolve_meta_sources(project_dir, clips_subdir)
    rows: list[tuple[str, str, str, str]] = []
    for mp in meta_files:
        try:
            meta = json.loads(mp.read_text(encoding="utf-8"))
        except Exception:
            continue
        stem = mp.name.replace("_meta.json", "")
        fn = f"{stem}.mp4"
        title = str(meta.get("title") or "").strip()
        desc = str(meta.get("description") or "").strip()
        tags_s = ""
        if include_tags:
            tags = meta.get("tags")
            if isinstance(tags, list):
                tags_s = ", ".join(str(t).strip() for t in tags if str(t).strip())
        rows.append((fn, title, desc, tags_s))
    return proj, rows


def build_metadata_markdown(
    project_dir: Path,
    clips_subdir: str = "clips_final",
    *,
    include_tags: bool = False,
    title_lang_note: str | None = None,
) -> str:
    """
    Собрать Markdown: заголовок проекта + для каждого clip_XXX — titre + description из *_meta.json.
    """
    project_dir = Path(project_dir)
    proj, meta_files = _resolve_meta_sources(project_dir, clips_subdir)

    source_title = str(proj.get("source_title") or "").strip()
    lang_from_project = str(proj.get("language") or "").strip()

    lines: list[str] = []
    doc_title = source_title or project_dir.name
    lines.append(f"# Métadonnées des clips — {doc_title}")
    lines.append("")
    meta_note = title_lang_note or lang_from_project or ""
    if meta_note:
        lines.append(
            f"> Projet: `{project_dir.name}` · Langue des textes: **{meta_note}** · {len(meta_files)} clips"
        )
    else:
        lines.append(f"> Projet: `{project_dir.name}` · {len(meta_files)} clips")
    lines.append("")
    lines.append("---")
    lines.append("")

    for mp in meta_files:
        try:
            meta = json.loads(mp.read_text(encoding="utf-8"))
        except Exception:
            continue
        stem = mp.name.replace("_meta.json", "")
        title = str(meta.get("title") or "").strip()
        desc = str(meta.get("description") or "").strip()
        lines.append(f"## {stem}.mp4")
        lines.append("")
        lines.append("**Titre**")
        lines.append("")
        lines.append(title if title else "—")
        lines.append("")
        lines.append("**Description**")
        lines.append("")
        lines.append(desc if desc else "—")
        if include_tags:
            tags = meta.get("tags")
            if isinstance(tags, list) and tags:
                lines.append("")
                lines.append("**Tags**")
                lines.append("")
                lines.append(", ".join(str(t) for t in tags if str(t).strip()))
        lines.append("")
        lines.append("---")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def build_metadata_xlsx_bytes(
    project_dir: Path,
    clips_subdir: str = "clips_final",
    *,
    include_tags: bool = False,
) -> bytes:
    """Собрать .xlsx в память (колонки: Fichier, Titre, Description [, Tags])."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as e:
        raise ImportError(
            "Для экспорта в Excel установите: pip install openpyxl"
        ) from e

    project_dir = Path(project_dir)
    proj, rows = _iter_clip_rows(
        project_dir, clips_subdir, include_tags=include_tags
    )
    if not rows:
        raise FileNotFoundError("Нет строк для экспорта (пустые meta?)")

    wb = Workbook()
    ws = wb.active
    ws.title = "Clips"[:31]

    ncol = 4 if include_tags else 3
    source_title = str(proj.get("source_title") or project_dir.name).strip()
    lang = str(proj.get("language") or "").strip()
    head_bits = [str(project_dir.name), source_title]
    if lang:
        head_bits.append(lang)
    headline = " · ".join(head_bits)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c1 = ws.cell(row=1, column=1, value=headline)
    c1.font = Font(bold=True)
    c1.alignment = Alignment(wrap_text=True, vertical="top")

    headers = ["Fichier", "Titre", "Description"]
    if include_tags:
        headers.append("Tags")
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    wrap = Alignment(wrap_text=True, vertical="top")
    for fn, title, desc, tags_s in rows:
        if include_tags:
            ws.append([fn, title, desc, tags_s])
        else:
            ws.append([fn, title, desc])

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for cell in row:
            cell.alignment = wrap

    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 72
    if include_tags:
        ws.column_dimensions["D"].width = 36

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def write_metadata_doc(
    project_dir: Path,
    out_path: Path | None = None,
    clips_subdir: str = "clips_final",
    *,
    as_xlsx: bool = True,
    **kwargs,
) -> Path:
    """Сохранить metadata в проект: по умолчанию .xlsx, иначе .md (as_xlsx=False)."""
    project_dir = Path(project_dir)
    lang = ""
    pj = project_dir / "project.json"
    if pj.exists():
        try:
            lang = str(json.loads(pj.read_text(encoding="utf-8")).get("language") or "").strip()
        except Exception:
            pass
    suffix = f"_{lang}" if lang else ""

    if as_xlsx:
        data = build_metadata_xlsx_bytes(project_dir, clips_subdir, **kwargs)
        out = out_path or (project_dir / f"metadata_titles_descriptions{suffix}.xlsx")
        out = Path(out)
        out.write_bytes(data)
        return out

    text = build_metadata_markdown(project_dir, clips_subdir, **kwargs)
    out = out_path or (project_dir / f"metadata_titles_descriptions{suffix}.md")
    out = Path(out)
    out.write_text(text, encoding="utf-8")
    return out
